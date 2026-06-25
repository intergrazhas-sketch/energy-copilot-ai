"""Shared tabular file reader for CSV and Excel imports.

Reads an uploaded CSV / XLSX / XLS file into a list of string-valued row dicts
with normalized column names (trimmed, lower-cased, spaces -> underscore).
Keeps CSV behaviour identical to the previous implementations while adding
Excel support for the pilot, where forecast / actual / weather data arrives as
Excel files from the operating company and the weather station.
"""

import csv
import io
import os
from datetime import datetime

CSV_EXTENSIONS = (".csv",)
XLSX_EXTENSIONS = (".xlsx",)
XLS_EXTENSIONS = (".xls",)
SUPPORTED_EXTENSIONS = CSV_EXTENSIONS + XLSX_EXTENSIONS + XLS_EXTENSIONS


class TabularImportError(Exception):
    """Raised when an uploaded table cannot be read."""


def normalize_header(name: object) -> str:
    return str(name).strip().lower().replace(" ", "_")


def first_value(row: dict[str, str], aliases: tuple[str, ...]) -> str:
    """Return the first non-empty value among the given normalized aliases."""
    for alias in aliases:
        value = row.get(alias)
        if value is not None and value.strip():
            return value.strip()
    return ""


def has_any_column(fieldnames: set[str], aliases: tuple[str, ...]) -> bool:
    return bool(set(aliases) & fieldnames)


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_csv(raw_content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    try:
        text = raw_content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise TabularImportError("File must be UTF-8 encoded (CSV) or a valid Excel file") from exc

    text = text.strip()
    if not text:
        raise TabularImportError("File is empty")

    reader = csv.DictReader(io.StringIO(text))
    fieldnames = [normalize_header(name) for name in (reader.fieldnames or []) if name is not None]
    rows: list[dict[str, str]] = []
    for raw_row in reader:
        row: dict[str, str] = {}
        for key, value in raw_row.items():
            if key is None:
                continue
            row[normalize_header(key)] = value.strip() if isinstance(value, str) else _cell_to_str(value)
        rows.append(row)
    return fieldnames, rows


def _build_normalized_headers(header: list[object]) -> list[str | None]:
    headers: list[str | None] = []
    for cell in header:
        if cell is None or str(cell).strip() == "":
            headers.append(None)
        else:
            headers.append(normalize_header(cell))
    return headers


def _read_xlsx(raw_content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise TabularImportError("Excel (.xlsx) support is not installed on the server") from exc

    try:
        workbook = load_workbook(io.BytesIO(raw_content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises various errors
        raise TabularImportError("Could not read the Excel (.xlsx) file") from exc

    try:
        worksheet = workbook.worksheets[0]
        rows_iter = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration as exc:
            raise TabularImportError("File is empty") from exc

        norm_headers = _build_normalized_headers(list(header_row))
        fieldnames = [name for name in norm_headers if name]
        rows: list[dict[str, str]] = []
        for raw in rows_iter:
            if raw is None or all(cell is None for cell in raw):
                continue
            row: dict[str, str] = {}
            for index, name in enumerate(norm_headers):
                if not name:
                    continue
                value = raw[index] if index < len(raw) else None
                row[name] = _cell_to_str(value)
            rows.append(row)
    finally:
        workbook.close()

    return fieldnames, rows


def _read_xls(raw_content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise TabularImportError("Excel (.xls) support is not installed on the server") from exc

    try:
        book = xlrd.open_workbook(file_contents=raw_content)
        sheet = book.sheet_by_index(0)
    except Exception as exc:  # noqa: BLE001 - xlrd raises various errors
        raise TabularImportError("Could not read the Excel (.xls) file") from exc

    if sheet.nrows == 0:
        raise TabularImportError("File is empty")

    norm_headers = _build_normalized_headers(list(sheet.row_values(0)))
    fieldnames = [name for name in norm_headers if name]
    rows: list[dict[str, str]] = []
    for row_index in range(1, sheet.nrows):
        cells = sheet.row(row_index)
        if all(cell.ctype == xlrd.XL_CELL_EMPTY for cell in cells):
            continue
        row: dict[str, str] = {}
        for index, name in enumerate(norm_headers):
            if not name:
                continue
            if index >= len(cells):
                row[name] = ""
                continue
            cell = cells[index]
            if cell.ctype == xlrd.XL_CELL_DATE:
                row[name] = xlrd.xldate_as_datetime(cell.value, book.datemode).isoformat()
            else:
                row[name] = _cell_to_str(cell.value)
        rows.append(row)

    return fieldnames, rows


def read_tabular_upload(
    filename: str | None,
    raw_content: bytes,
) -> tuple[list[str], list[dict[str, str]]]:
    """Read a CSV/XLSX/XLS upload into (normalized_fieldnames, normalized_rows).

    Raises TabularImportError with a user-friendly message on failure.
    """
    extension = os.path.splitext(filename or "")[1].lower()
    if extension in XLSX_EXTENSIONS:
        return _read_xlsx(raw_content)
    if extension in XLS_EXTENSIONS:
        return _read_xls(raw_content)
    # Default to CSV (covers .csv and unknown/empty extensions, preserving prior behaviour).
    return _read_csv(raw_content)
