"""Detect weather/meteo station exports that must not break batch import."""

from __future__ import annotations

import io

from app.services.campbell_toa5 import is_campbell_toa5

WEATHER_SKIP_REASON = "weather_file_detected_not_imported_yet"

WEATHER_FILENAME_MARKERS = (
    "meteo",
    "weather",
    "метео",
    "toa5",
    "campbell",
    "slrw",
    "irradiance",
    "piran",
    "станция",
    "станц",
)

WEATHER_SHEET_MARKERS = (
    "meteo",
    "weather",
    "метео",
    "toa5",
    "irradiance",
    "piran",
    "slrw",
)

GENERATION_SHEET_MARKERS = ("генераци", "generation", "свод")


def _filename_suggests_weather(filename: str | None) -> bool:
    lowered = (filename or "").lower().replace(" ", "_")
    return any(marker in lowered for marker in WEATHER_FILENAME_MARKERS)


def _sheet_name_suggests_weather(sheet_name: str) -> bool:
    lowered = sheet_name.lower().replace(" ", "_")
    return any(marker in lowered for marker in WEATHER_SHEET_MARKERS)


def _sheet_name_suggests_generation(sheet_name: str) -> bool:
    lowered = sheet_name.lower().replace(" ", "_")
    return any(marker in lowered for marker in GENERATION_SHEET_MARKERS)


def _header_row_suggests_weather(cells: list[object]) -> bool:
    joined = " ".join(str(cell).strip().lower().replace(" ", "_") for cell in cells if cell is not None)
    if not joined:
        return False
    weather_hits = sum(1 for marker in WEATHER_SHEET_MARKERS if marker in joined)
    generation_hits = sum(1 for marker in ("сумма_генерации", "прогноз_генераци", "generation") if marker in joined)
    return weather_hits >= 2 and generation_hits == 0


def _workbook_is_weather_only_xlsx(raw: bytes) -> bool:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return False

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        return False

    try:
        sheet_names = list(workbook.sheetnames)
        if not sheet_names:
            return False

        has_generation_sheet = any(_sheet_name_suggests_generation(name) for name in sheet_names)
        if has_generation_sheet:
            return False

        if any(_sheet_name_suggests_weather(name) for name in sheet_names):
            return True

        first_sheet = workbook[sheet_names[0]]
        header_row: list[object] = []
        for index, row in enumerate(first_sheet.iter_rows(values_only=True)):
            header_row = list(row)
            if index == 0:
                break
        return _header_row_suggests_weather(header_row)
    except Exception:
        return False
    finally:
        workbook.close()


def is_weather_meteo_file(filename: str | None, raw: bytes) -> bool:
    """Return True when the upload is a meteo/weather export, not generation."""
    if is_campbell_toa5(filename, raw):
        return True
    if _filename_suggests_weather(filename):
        return True
    lowered = (filename or "").lower()
    if lowered.endswith((".xlsx", ".xlsm")):
        return _workbook_is_weather_only_xlsx(raw)
    return False
